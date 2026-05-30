from __future__ import annotations

from dataclasses import asdict
from datetime import date, datetime
from pathlib import Path
from typing import Any
from uuid import uuid4

import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.modules.context_pack import context_pack_identity
from app.modules.field_mapping import build_field_profile
from app.modules.schemas import (
    ColumnSummary,
    DataSourceRef,
    PreviewData,
    SchemaSummary,
    TableData,
    WorkbookSheet,
)


class IngestionError(Exception):
    def __init__(self, code: str, message: str, details: dict[str, Any] | None = None) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.details = details or {}


SUPPORTED_SUFFIXES = {".csv", ".xlsx"}


def load_table_from_path(
    path: Path,
    *,
    original_filename: str | None = None,
    selected_sheet: str | None = None,
    source_id: str | None = None,
) -> TableData:
    filename = original_filename or path.name
    suffix = Path(filename).suffix.lower()

    if suffix not in SUPPORTED_SUFFIXES:
        raise IngestionError(
            "UNSUPPORTED_FILE_TYPE",
            "当前仅支持 CSV 和 Excel（.xlsx），请另存为 .xlsx 或 CSV 后重试。",
            {"filename": filename},
        )

    if suffix == ".csv":
        dataframe = read_csv(path)
        data_source_type = "csv"
        workbook_sheets: list[WorkbookSheet] = []
        selected = None
    else:
        dataframe, workbook_sheets, selected = read_xlsx(path, selected_sheet=selected_sheet)
        data_source_type = "xlsx"

    dataframe = normalize_dataframe(dataframe)
    schema = build_schema_summary(dataframe)
    preview = build_preview(dataframe)
    source_ref = DataSourceRef(
        id=source_id or f"{data_source_type}-{uuid4().hex[:12]}",
        type=data_source_type,
        name=filename,
        location=str(path),
        selected_sheet=selected,
    )

    return TableData(
        data_source_ref=source_ref,
        dataframe=dataframe,
        schema_summary=schema,
        preview=preview,
        workbook_sheets=workbook_sheets,
        selected_sheet=selected,
    )


def read_csv(path: Path) -> pd.DataFrame:
    try:
        return pd.read_csv(path, encoding="utf-8")
    except UnicodeDecodeError as exc:
        raise IngestionError("CSV_PARSE_FAILED", "无法解析这份 CSV，可能是编码问题。") from exc
    except pd.errors.ParserError as exc:
        raise IngestionError(
            "CSV_PARSE_FAILED",
            "无法解析这份 CSV，可能是分隔符或表头问题。",
        ) from exc


def read_xlsx(
    path: Path,
    *,
    selected_sheet: str | None = None,
) -> tuple[pd.DataFrame, list[WorkbookSheet], str]:
    try:
        workbook = load_workbook(path, data_only=False)
    except Exception as exc:
        raise IngestionError(
            "XLSX_PARSE_FAILED",
            "无法读取这份 Excel，请另存为 .xlsx 后重试。",
        ) from exc

    sheets = build_sheet_metadata(workbook.worksheets)
    visible_non_empty = [sheet.name for sheet in sheets if sheet.visible and not sheet.empty]
    if not visible_non_empty:
        raise IngestionError("XLSX_NO_VISIBLE_SHEET", "未找到可分析的非空可见工作表。")

    sheet_name = selected_sheet or visible_non_empty[0]
    if sheet_name not in visible_non_empty:
        raise IngestionError(
            "XLSX_NO_VISIBLE_SHEET",
            "选择的工作表不存在、被隐藏或为空。",
            {"selected_sheet": sheet_name},
        )

    ensure_formula_cache_available(path, sheet_name)

    try:
        dataframe = pd.read_excel(path, sheet_name=sheet_name, engine="openpyxl")
    except Exception as exc:
        raise IngestionError(
            "XLSX_PARSE_FAILED",
            "无法读取这份 Excel，请检查表头或保存后重试。",
        ) from exc

    selected_sheets = [
        WorkbookSheet(sheet.name, sheet.visible, sheet.empty, selected=sheet.name == sheet_name)
        for sheet in sheets
        if sheet.visible
    ]
    return dataframe, selected_sheets, sheet_name


def build_sheet_metadata(worksheets: list[Worksheet]) -> list[WorkbookSheet]:
    sheets: list[WorkbookSheet] = []
    for worksheet in worksheets:
        sheets.append(
            WorkbookSheet(
                name=worksheet.title,
                visible=worksheet.sheet_state == "visible",
                empty=is_sheet_empty(worksheet),
            )
        )
    return sheets


def is_sheet_empty(worksheet: Worksheet) -> bool:
    for row in worksheet.iter_rows(values_only=True):
        if any(value not in (None, "") for value in row):
            return False
    return True


def ensure_formula_cache_available(path: Path, sheet_name: str) -> None:
    formula_workbook = load_workbook(path, data_only=False)
    value_workbook = load_workbook(path, data_only=True)
    formula_sheet = formula_workbook[sheet_name]
    value_sheet = value_workbook[sheet_name]

    for row in formula_sheet.iter_rows():
        for cell in row:
            if cell.data_type == "f" and value_sheet[cell.coordinate].value is None:
                raise IngestionError(
                    "XLSX_PARSE_FAILED",
                    "Excel 公式缺少缓存值，请用 Excel 打开并保存后重试。",
                    {"sheet": sheet_name, "cell": cell.coordinate},
                )


def normalize_dataframe(dataframe: pd.DataFrame) -> pd.DataFrame:
    dataframe = dataframe.dropna(axis=0, how="all").dropna(axis=1, how="all").copy()
    dataframe.columns = [str(column).strip() for column in dataframe.columns]
    if dataframe.empty or not dataframe.columns.any():
        raise IngestionError("CSV_PARSE_FAILED", "数据为空或缺少表头。")
    unsupported_columns = [
        column for column in dataframe.columns if column.lower().startswith("unnamed:")
    ]
    if len(unsupported_columns) == len(dataframe.columns):
        raise IngestionError(
            "XLSX_UNSUPPORTED_FEATURE",
            "当前版本需要首行是清晰表头的普通二维表格。",
        )
    if (
        has_complex_header(dataframe, unsupported_columns)
        or has_data_like_header(dataframe.columns)
    ):
        raise IngestionError(
            "XLSX_UNSUPPORTED_FEATURE",
            "当前版本需要首行是清晰表头的普通二维表格。",
        )
    return dataframe


def has_complex_header(dataframe: pd.DataFrame, unsupported_columns: list[str]) -> bool:
    if not unsupported_columns:
        return False
    supported_columns = len(dataframe.columns) - len(unsupported_columns)
    if supported_columns > 1:
        return False
    first_row = dataframe.iloc[0] if len(dataframe.index) else None
    if first_row is None:
        return True
    meaningful_values = [value for value in first_row.tolist() if value not in (None, "")]
    return len(meaningful_values) >= 2


def has_data_like_header(columns: pd.Index) -> bool:
    labels = [str(column).strip() for column in columns]
    if len(labels) < 2:
        return False
    date_like_count = sum(is_date_like_label(label) for label in labels)
    numeric_like_count = sum(is_numeric_like_label(label) for label in labels)
    return date_like_count >= 1 and numeric_like_count >= 1


def is_date_like_label(label: str) -> bool:
    if not label:
        return False
    parsed = pd.to_datetime([label], errors="coerce", format="mixed")
    return bool(parsed.notna()[0])


def is_numeric_like_label(label: str) -> bool:
    if not label:
        return False
    try:
        float(label.replace(",", ""))
    except ValueError:
        return False
    return True


def build_schema_summary(dataframe: pd.DataFrame) -> SchemaSummary:
    columns = []
    for column in dataframe.columns:
        data_type = infer_series_type(dataframe[column])
        min_value, max_value = column_min_max(dataframe[column], data_type)
        columns.append(
            ColumnSummary(
                name=str(column),
                data_type=data_type,
                nullable=bool(dataframe[column].isna().any()),
                sample_values=sample_values(dataframe[column]),
                distinct_count=int(dataframe[column].nunique(dropna=True)),
                min_value=min_value,
                max_value=max_value,
            )
        )
    return SchemaSummary(
        tables=[
            {
                "name": "sales_orders",
                "columns": [asdict(column) for column in columns],
                "row_count_estimate": int(len(dataframe.index)),
            }
        ],
        columns=columns,
        row_count_estimate=int(len(dataframe.index)),
    )


def infer_series_type(series: pd.Series) -> str:
    non_null = series.dropna()
    if non_null.empty:
        return "string"
    if pd.api.types.is_datetime64_any_dtype(series):
        return "date"
    if pd.api.types.is_numeric_dtype(series):
        return "number"
    parsed = pd.to_datetime(non_null.head(20), errors="coerce", format="mixed")
    if parsed.notna().mean() >= 0.8:
        return "date"
    if non_null.nunique(dropna=True) <= max(20, int(len(non_null) * 0.2)):
        return "category"
    return "string"


def column_min_max(series: pd.Series, data_type: str) -> tuple[Any, Any]:
    """对 date / number 列计算 min/max，供 LLM 生成 SQL 时圈定时间/数值窗口。"""
    if data_type not in {"date", "number"}:
        return None, None
    non_null = series.dropna()
    if non_null.empty:
        return None, None
    if data_type == "date":
        parsed = pd.to_datetime(non_null, errors="coerce")
        parsed = parsed.dropna()
        if parsed.empty:
            return None, None
        return parsed.min().date().isoformat(), parsed.max().date().isoformat()
    numeric = pd.to_numeric(non_null, errors="coerce").dropna()
    if numeric.empty:
        return None, None
    return float(numeric.min()), float(numeric.max())


def sample_values(series: pd.Series) -> list[Any]:
    values = []
    for value in series.dropna().head(5).tolist():
        values.append(normalize_value(value))
    return values


def build_preview(dataframe: pd.DataFrame) -> PreviewData:
    head = dataframe.head(20)
    tail = dataframe.tail(5)
    return PreviewData(head=records(head), tail=records(tail))


def records(dataframe: pd.DataFrame) -> list[dict[str, Any]]:
    return [
        {str(key): normalize_value(value) for key, value in row.items()}
        for row in dataframe.to_dict(orient="records")
    ]


def normalize_value(value: Any) -> Any:
    if pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.date().isoformat()
    if isinstance(value, datetime | date):
        return value.isoformat()
    return value


def table_to_response(table: TableData) -> dict[str, Any]:
    profile = build_field_profile(table.schema_summary)
    return {
        **context_pack_identity(),
        "data_source_ref": asdict(table.data_source_ref),
        "row_count": table.row_count,
        "column_count": table.column_count,
        "schema_summary": asdict(table.schema_summary),
        "preview": asdict(table.preview),
        "workbook_sheets": [asdict(sheet) for sheet in table.workbook_sheets],
        "selected_sheet": table.selected_sheet,
        "field_profile": asdict(profile),
    }
