from pathlib import Path

import pytest

from app.modules.field_mapping import build_field_profile
from app.modules.file_ingestion import IngestionError, load_table_from_path


def test_csv_file_returns_preview_and_field_profile(tmp_path: Path):
    source = tmp_path / "orders.csv"
    source.write_text(
        "\n".join(
            [
                "order_date,net_sales_amount,store_name,category_l1,channel",
                "2026-05-11,1000,上海徐汇旗舰店,童装,线下",
                "2026-05-12,800,广州天河店,服装,抖音",
            ]
        ),
        encoding="utf-8",
    )

    table = load_table_from_path(source, original_filename="orders.csv")
    profile = build_field_profile(table.schema_summary)

    assert table.data_source_ref.type == "csv"
    assert table.row_count == 2
    assert table.preview.head[0]["store_name"] == "上海徐汇旗舰店"
    assert profile.is_valid is True
    assert profile.mappings["net_sales_amount"].canonical_field == "net_sales_amount"
    assert profile.mappings["order_date"].canonical_field == "order_date"

    date_col = next(c for c in table.schema_summary.columns if c.name == "order_date")
    sales_col = next(c for c in table.schema_summary.columns if c.name == "net_sales_amount")
    store_col = next(c for c in table.schema_summary.columns if c.name == "store_name")
    assert date_col.min_value == "2026-05-11"
    assert date_col.max_value == "2026-05-12"
    assert sales_col.min_value == 800.0
    assert sales_col.max_value == 1000.0
    assert store_col.min_value is None and store_col.max_value is None


def test_xlsx_defaults_to_first_non_empty_visible_sheet_and_can_switch(tmp_path: Path):
    from openpyxl import Workbook

    source = tmp_path / "orders.xlsx"
    workbook = Workbook()
    empty = workbook.active
    empty.title = "empty"
    hidden = workbook.create_sheet("hidden-data")
    hidden.sheet_state = "hidden"
    hidden.append(["order_date", "net_sales_amount", "store_name"])
    hidden.append(["2026-05-11", 1, "隐藏门店"])
    may = workbook.create_sheet("may-orders")
    may.append(["order_date", "net_sales_amount", "store_name", "category_l1", "channel"])
    may.append(["2026-05-11", 1200, "上海徐汇旗舰店", "童装", "线下"])
    june = workbook.create_sheet("june-orders")
    june.append(["订单日期", "销售额", "门店", "类目", "渠道"])
    june.append(["2026-06-01", 500, "广州天河店", "服装", "抖音"])
    workbook.save(source)

    default_table = load_table_from_path(source, original_filename="orders.xlsx")
    switched_table = load_table_from_path(
        source,
        original_filename="orders.xlsx",
        selected_sheet="june-orders",
    )
    switched_profile = build_field_profile(switched_table.schema_summary)

    assert default_table.data_source_ref.type == "xlsx"
    assert default_table.selected_sheet == "may-orders"
    assert [sheet.name for sheet in default_table.workbook_sheets] == [
        "empty",
        "may-orders",
        "june-orders",
    ]
    assert switched_table.selected_sheet == "june-orders"
    assert switched_table.preview.head[0]["门店"] == "广州天河店"
    assert switched_profile.is_valid is True
    assert switched_profile.mappings["销售额"].canonical_field == "net_sales_amount"
    assert switched_profile.mappings["订单日期"].canonical_field == "order_date"


def test_empty_xlsx_workbook_is_rejected(tmp_path: Path):
    from openpyxl import Workbook

    source = tmp_path / "empty.xlsx"
    workbook = Workbook()
    workbook.save(source)

    with pytest.raises(IngestionError) as exc:
        load_table_from_path(source, original_filename="empty.xlsx")

    assert exc.value.code == "XLSX_NO_VISIBLE_SHEET"


def test_xlsx_with_complex_header_is_rejected(tmp_path: Path):
    from openpyxl import Workbook

    source = tmp_path / "complex-header.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "orders"
    sheet.append(["经营日报", None, None, None])
    sheet.append(["订单日期", "销售额", "门店", "渠道"])
    sheet.append(["2026-05-11", 1000, "上海徐汇旗舰店", "线下"])
    workbook.save(source)

    with pytest.raises(IngestionError) as exc:
        load_table_from_path(source, original_filename="complex-header.xlsx")

    assert exc.value.code == "XLSX_UNSUPPORTED_FEATURE"


def test_xlsx_without_header_row_is_rejected(tmp_path: Path):
    from openpyxl import Workbook

    source = tmp_path / "no-header.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "orders"
    sheet.append(["2026-05-11", 1000, "上海徐汇旗舰店", "线下"])
    sheet.append(["2026-05-12", 800, "广州天河店", "抖音"])
    workbook.save(source)

    with pytest.raises(IngestionError) as exc:
        load_table_from_path(source, original_filename="no-header.xlsx")

    assert exc.value.code == "XLSX_UNSUPPORTED_FEATURE"


def test_xlsx_formula_without_cached_value_is_rejected(tmp_path: Path):
    from openpyxl import Workbook

    source = tmp_path / "formula.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "orders"
    sheet.append(["order_date", "net_sales_amount", "store_name"])
    sheet.append(["2026-05-11", "=500+500", "上海徐汇旗舰店"])
    workbook.save(source)

    with pytest.raises(IngestionError) as exc:
        load_table_from_path(source, original_filename="formula.xlsx")

    assert exc.value.code == "XLSX_PARSE_FAILED"
    assert exc.value.details == {"sheet": "orders", "cell": "B2"}


def test_unsupported_legacy_excel_type_is_rejected(tmp_path: Path):
    source = tmp_path / "orders.xls"
    source.write_bytes(b"legacy")

    with pytest.raises(IngestionError) as exc:
        load_table_from_path(source, original_filename="orders.xls")

    assert exc.value.code == "UNSUPPORTED_FILE_TYPE"


def test_missing_key_fields_mark_profile_invalid(tmp_path: Path):
    source = tmp_path / "orders.csv"
    source.write_text("store_name,category_l1\n上海徐汇旗舰店,童装\n", encoding="utf-8")

    table = load_table_from_path(source, original_filename="orders.csv")
    profile = build_field_profile(table.schema_summary)

    assert profile.is_valid is False
    assert profile.missing_key_fields == ["net_sales_amount", "order_date"]
